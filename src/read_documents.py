"""Read text from supplied DOCX files and scanned image files.

DOCX files are read locally with python-docx. PNG images are passed to the
locally installed Tesseract OCR engine through pytesseract. This module only
turns pages into text; it does not interpret agreement clauses.
"""
from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from PIL import Image
import pytesseract

IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}
SUPPORTED_SUFFIXES = {".docx", *IMAGE_SUFFIXES}


def read_docx(path: Path) -> str:
    """Read paragraphs and tables from a Word document."""
    document = Document(path)
    parts: list[str] = []
    parts.extend(p.text for p in document.paragraphs if p.text.strip())
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def read_image(path: Path) -> str:
    """Use the local Tesseract installation to recognize a scanned image."""
    with Image.open(path) as image:
        return pytesseract.image_to_string(image, lang="eng").strip()


def extract_text(path: Path) -> str:
    """Choose a reader from the actual file extension and return its text."""
    if not path.is_file():
        raise FileNotFoundError(f"Document does not exist: {path}")
    suffix = path.suffix.lower()
    if suffix == ".docx":
        text = read_docx(path)
    elif suffix in IMAGE_SUFFIXES:
        text = read_image(path)
    else:
        raise ValueError(
            f"Unsupported file type {suffix!r} for {path.name}. "
            f"Supported types: {', '.join(sorted(SUPPORTED_SUFFIXES))}."
        )
    if not text.strip():
        raise ValueError(
            f"No text could be read from {path.name}. For PNG images, check "
            "that Tesseract OCR is installed and the scan is legible."
        )
    return text


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Read DOCX agreements and OCR scanned image agreements."
    )
    parser.add_argument("path", type=Path, help="A document or folder of documents")
    parser.add_argument(
        "--show-text",
        action="store_true",
        help="Print recognized agreement text (may contain sensitive information)",
    )
    args = parser.parse_args()

    paths = [args.path] if args.path.is_file() else sorted(args.path.rglob("*"))
    documents = [p for p in paths if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES]
    if not documents:
        parser.error(f"No supported DOCX or PNG files found at {args.path}")

    failures = 0
    for path in documents:
        try:
            text = extract_text(path)
        except Exception as exc:
            print(f"ERROR {path.name}: {exc}")
            failures += 1
            continue
        print(f"{path.name}: extracted {len(text)} characters")
        if args.show_text:
            print(text)
            print("-" * 40)
    if failures:
        raise SystemExit(f"Could not read {failures} document(s).")


if __name__ == "__main__":
    main()
